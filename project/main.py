from flask import Flask, render_template, request, redirect, url_for  # type: ignore
from openpyxl import Workbook, load_workbook
import os

app = Flask(__name__)

# Hardcoded credentials
USERNAME = "admin"
PASSWORD = "1234"

# Excel file path
EXCEL_FILE = 'farmer_data.xlsx'

# Initialize Excel file if it doesn't exist
if not os.path.exists(EXCEL_FILE):
    wb = Workbook()
    ws = wb.active
    ws.title = "Farmers"
    ws.append(['Name', 'Village', 'Phone', 'Main Crops'])
    wb.save(EXCEL_FILE)

@app.route('/', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form['username']
        password = request.form['password']
        if username == USERNAME and password == PASSWORD:
            return redirect(url_for('home'))
        else:
            return "Invalid Credentials. Try again."
    return render_template('login.html')

@app.route('/home')
def home():
    return render_template('home.html')

@app.route('/register', methods=['GET', 'POST'])
def register():
    if request.method == 'POST':
        name = request.form['name']
        village = request.form['village']
        phone = request.form['phone']
        crops = request.form['crops']

        wb = load_workbook(EXCEL_FILE)
        ws = wb.active
        ws.append([name, village, phone, crops])
        wb.save(EXCEL_FILE)

        return redirect(url_for('home'))

    return render_template('register.html')

@app.route('/booking', methods=['GET', 'POST'])
def booking():
    if request.method == 'POST':
        try:
            booking_name = request.form['name']
            booking_village = request.form['village']
            booking_phone = request.form['phone']
            booking_date = request.form['date']
            booking_time = request.form['time']
        except KeyError as e:
            return f"Missing form field: {e}"

        wb = load_workbook(EXCEL_FILE)

        if "Bookings" not in wb.sheetnames:
            wb.create_sheet("Bookings")
            booking_ws = wb["Bookings"]
            booking_ws.append(['Name', 'Village', 'Phone', 'Booking Date', 'Booking Time'])
        else:
            booking_ws = wb["Bookings"]

        booking_ws.append([booking_name, booking_village, booking_phone, booking_date, booking_time])
        wb.save(EXCEL_FILE)

        return redirect(url_for('home'))

    return render_template('booking.html')

@app.route('/shares')
def shares():
    return render_template('shares.html')

@app.route('/project')
def project():
    return render_template('project.html')

@app.route('/contact', methods=['GET', 'POST'])
def contact():
    if request.method == 'POST':
        try:
            contact_name = request.form['name']
            contact_email = request.form['email']
            contact_message = request.form['message']
        except KeyError as e:
            return f"Missing form field: {e}"

        wb = load_workbook(EXCEL_FILE)

        if "Contacts" not in wb.sheetnames:
            wb.create_sheet("Contacts")
            contact_ws = wb["Contacts"]
            contact_ws.append(['Name', 'Email', 'Message'])
        else:
            contact_ws = wb["Contacts"]

        contact_ws.append([contact_name, contact_email, contact_message])
        wb.save(EXCEL_FILE)

        return redirect(url_for('home'))

    return render_template('contact.html')

if __name__ == "__main__":
    app.run(debug=True)
